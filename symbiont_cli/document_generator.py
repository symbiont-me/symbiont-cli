import os
from datetime import datetime
from pathlib import Path
from typing import List, Optional, Dict, Any

from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib import colors


def filter_metadata(metadata: dict) -> dict:
    """Filter out unnecessary metadata fields for cleaner output"""
    # Fields to exclude from metadata display
    excluded_fields = {
        'producer', 'creator', 'author', 'subject', 'keywords',
        'creation_date', 'modification_date', 'trapped', 'encrypted'
    }
    
    # Only keep relevant fields
    relevant_fields = {
        'source', 'title', 'page', 'relevance_score', 'file_path'
    }
    
    filtered = {}
    for key, value in metadata.items():
        # Convert key to lowercase for case-insensitive comparison
        key_lower = key.lower()
        
        # Include if it's a relevant field and not in excluded list
        if key_lower in relevant_fields or (key_lower not in excluded_fields and key in relevant_fields):
            filtered[key] = value
            
    return filtered


class DocumentData:
    """Data structure to hold all information needed for document generation"""
    
    def __init__(self):
        self.query: str = ""
        self.collection_name: str = ""
        self.llm_response: Optional[str] = None
        self.search_results: List[Dict[str, Any]] = []
        self.config_info: Dict[str, Any] = {}
        self.processing_info: Dict[str, Any] = {}
        self.timestamp: datetime = datetime.now()


class DocumentGenerator:
    """Generates clean, professional documents from search results and LLM responses"""
    
    def __init__(self):
        self.data = DocumentData()
    
    def set_data(self, query: str, collection_name: str, search_results: List[Any], 
                 llm_response: Optional[str] = None, config_info: Optional[Dict] = None,
                 processing_info: Optional[Dict] = None):
        """Set the data for document generation"""
        self.data.query = query
        self.data.collection_name = collection_name
        self.data.llm_response = llm_response
        self.data.config_info = config_info or {}
        self.data.processing_info = processing_info or {}
        
        # Convert search results to standardized format
        self.data.search_results = []
        for result in search_results:
            result_data = {
                'content': getattr(result, 'page_content', str(result)),
                'metadata': getattr(result, 'metadata', {}),
                'relevance_score': getattr(result, 'metadata', {}).get('relevance_score', 'N/A')
            }
            self.data.search_results.append(result_data)
    
    def generate_document(self, output_path: str):
        """Generate document based on file extension"""
        file_path = Path(output_path)
        
        if file_path.suffix.lower() == '.docx':
            self._generate_word_document(output_path)
        elif file_path.suffix.lower() == '.pdf':
            self._generate_pdf_document(output_path)
        elif file_path.suffix.lower() == '.html':
            self._generate_html_document(output_path)
        else:
            raise ValueError(f"Unsupported file format: {file_path.suffix}")
    
    def _generate_word_document(self, output_path: str):
        """Generate a Word document"""
        doc = Document()
        
        # Title Page
        title = doc.add_heading('Research Query Results', 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph()
        
        # Query Information
        doc.add_heading('Query Information', level=1)
        
        info_table = doc.add_table(rows=4, cols=2)
        info_table.style = 'Table Grid'
        
        info_table.cell(0, 0).text = 'Query:'
        info_table.cell(0, 1).text = self.data.query
        
        info_table.cell(1, 0).text = 'Collection:'
        info_table.cell(1, 1).text = self.data.collection_name
        
        info_table.cell(2, 0).text = 'Generated:'
        info_table.cell(2, 1).text = self.data.timestamp.strftime('%Y-%m-%d %H:%M:%S')
        
        info_table.cell(3, 0).text = 'Results Found:'
        info_table.cell(3, 1).text = str(len(self.data.search_results))
        
        doc.add_paragraph()
        
        # LLM Response Section
        if self.data.llm_response:
            doc.add_heading('AI Generated Summary', level=1)
            response_para = doc.add_paragraph(self.data.llm_response)
            response_para.style = 'Intense Quote'
            doc.add_paragraph()
        
        # Configuration Information
        if self.data.config_info:
            doc.add_heading('Configuration', level=1)
            for key, value in self.data.config_info.items():
                if isinstance(value, dict):
                    doc.add_paragraph(f"{key.title()}:", style='Heading 2')
                    for sub_key, sub_value in value.items():
                        doc.add_paragraph(f"  • {sub_key}: {sub_value}")
                else:
                    doc.add_paragraph(f"{key.title()}: {value}")
            doc.add_paragraph()
        
        # Search Results
        if self.data.search_results:
            doc.add_heading('Retrieved Documents', level=1)
            
            for i, result in enumerate(self.data.search_results, 1):
                doc.add_heading(f'Document {i}', level=2)
                
                # Metadata table
                if result['metadata']:
                    filtered_metadata = filter_metadata(result['metadata'])
                    if filtered_metadata:
                        metadata_table = doc.add_table(rows=len(filtered_metadata), cols=2)
                        metadata_table.style = 'Light Shading'
                        
                        for row, (key, value) in enumerate(filtered_metadata.items()):
                            metadata_table.cell(row, 0).text = key.title()
                            metadata_table.cell(row, 1).text = str(value)
                
                # Content
                doc.add_paragraph('Content:', style='Heading 3')
                content_para = doc.add_paragraph(result['content'])
                content_para.style = 'Body Text'
                
                if i < len(self.data.search_results):
                    doc.add_page_break()
        
        # Save document
        doc.save(output_path)
    
    def _generate_pdf_document(self, output_path: str):
        """Generate a PDF document"""
        doc = SimpleDocTemplate(output_path, pagesize=letter,
                              rightMargin=72, leftMargin=72,
                              topMargin=72, bottomMargin=18)
        
        # Get styles
        styles = getSampleStyleSheet()
        story = []
        
        # Custom styles
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            spaceAfter=30,
            alignment=1  # Center
        )
        
        header_style = ParagraphStyle(
            'CustomHeader',
            parent=styles['Heading2'],
            fontSize=16,
            spaceAfter=12,
            textColor=colors.darkblue
        )
        
        # Title
        story.append(Paragraph("Research Query Results", title_style))
        story.append(Spacer(1, 20))
        
        # Query Information
        story.append(Paragraph("Query Information", header_style))
        
        query_data = [
            ['Query:', self.data.query],
            ['Collection:', self.data.collection_name],
            ['Generated:', self.data.timestamp.strftime('%Y-%m-%d %H:%M:%S')],
            ['Results Found:', str(len(self.data.search_results))]
        ]
        
        query_table = Table(query_data, colWidths=[1.5*inch, 4*inch])
        query_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (0, -1), colors.lightgrey),
            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
            ('BACKGROUND', (1, 0), (1, -1), colors.beige),
            ('GRID', (0, 0), (-1, -1), 1, colors.black)
        ]))
        
        story.append(query_table)
        story.append(Spacer(1, 20))
        
        # LLM Response
        if self.data.llm_response:
            story.append(Paragraph("AI Generated Summary", header_style))
            story.append(Paragraph(self.data.llm_response, styles['BodyText']))
            story.append(Spacer(1, 20))
        
        # Configuration
        if self.data.config_info:
            story.append(Paragraph("Configuration", header_style))
            for key, value in self.data.config_info.items():
                if isinstance(value, dict):
                    story.append(Paragraph(f"<b>{key.title()}:</b>", styles['BodyText']))
                    for sub_key, sub_value in value.items():
                        story.append(Paragraph(f"  • {sub_key}: {sub_value}", styles['BodyText']))
                else:
                    story.append(Paragraph(f"<b>{key.title()}:</b> {value}", styles['BodyText']))
            story.append(Spacer(1, 20))
        
        # Search Results
        if self.data.search_results:
            story.append(Paragraph("Retrieved Documents", header_style))
            
            for i, result in enumerate(self.data.search_results, 1):
                # Document header
                doc_header = ParagraphStyle(
                    'DocumentHeader',
                    parent=styles['Heading3'],
                    fontSize=14,
                    spaceAfter=10,
                    textColor=colors.darkgreen
                )
                story.append(Paragraph(f"Document {i}", doc_header))
                
                # Metadata
                if result['metadata']:
                    filtered_metadata = filter_metadata(result['metadata'])
                    if filtered_metadata:
                        metadata_rows = [[key.title(), str(value)] for key, value in filtered_metadata.items()]
                        metadata_table = Table(metadata_rows, colWidths=[1.5*inch, 4*inch])
                        metadata_table.setStyle(TableStyle([
                            ('BACKGROUND', (0, 0), (0, -1), colors.lightblue),
                            ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
                            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
                            ('FONTSIZE', (0, 0), (-1, -1), 9),
                            ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
                            ('GRID', (0, 0), (-1, -1), 1, colors.black)
                        ]))
                        story.append(metadata_table)
                        story.append(Spacer(1, 10))
                
                # Content
                story.append(Paragraph("<b>Content:</b>", styles['BodyText']))
                # Split long content into smaller paragraphs
                content = result['content']
                if len(content) > 500:
                    # Split into chunks for better formatting
                    words = content.split()
                    chunks = []
                    current_chunk = []
                    current_length = 0
                    
                    for word in words:
                        if current_length + len(word) > 500 and current_chunk:
                            chunks.append(' '.join(current_chunk))
                            current_chunk = [word]
                            current_length = len(word)
                        else:
                            current_chunk.append(word)
                            current_length += len(word) + 1
                    
                    if current_chunk:
                        chunks.append(' '.join(current_chunk))
                    
                    for chunk in chunks:
                        story.append(Paragraph(chunk, styles['BodyText']))
                else:
                    story.append(Paragraph(content, styles['BodyText']))
                
                if i < len(self.data.search_results):
                    story.append(Spacer(1, 30))
        
        # Build PDF
        doc.build(story)
    
    def _generate_html_document(self, output_path: str):
        """Generate an HTML document"""
        html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Research Query Results</title>
    <style>
        body {{
            font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
            line-height: 1.6;
            margin: 0;
            padding: 20px;
            background-color: #f5f5f5;
            color: #333;
        }}
        .container {{
            max-width: 1200px;
            margin: 0 auto;
            background-color: white;
            padding: 30px;
            border-radius: 8px;
            box-shadow: 0 2px 10px rgba(0,0,0,0.1);
        }}
        h1 {{
            color: #2c3e50;
            text-align: center;
            border-bottom: 3px solid #3498db;
            padding-bottom: 10px;
            margin-bottom: 30px;
        }}
        h2 {{
            color: #34495e;
            border-left: 4px solid #3498db;
            padding-left: 15px;
            margin-top: 30px;
        }}
        h3 {{
            color: #27ae60;
            margin-top: 25px;
        }}
        .query-info {{
            background-color: #ecf0f1;
            padding: 20px;
            border-radius: 5px;
            margin-bottom: 25px;
        }}
        .query-info table {{
            width: 100%;
            border-collapse: collapse;
        }}
        .query-info td {{
            padding: 8px 12px;
            border: 1px solid #bdc3c7;
        }}
        .query-info td:first-child {{
            background-color: #d5dbdb;
            font-weight: bold;
            width: 150px;
        }}
        .llm-response {{
            background-color: #e8f5e8;
            border-left: 4px solid #27ae60;
            padding: 20px;
            margin: 20px 0;
            border-radius: 0 5px 5px 0;
            font-style: italic;
        }}
        .config-section {{
            background-color: #fdf2e9;
            padding: 15px;
            border-radius: 5px;
            margin: 20px 0;
        }}
        .config-section ul {{
            margin: 10px 0;
            padding-left: 20px;
        }}
        .document {{
            border: 1px solid #ddd;
            margin: 20px 0;
            border-radius: 5px;
            overflow: hidden;
        }}
        .document-header {{
            background-color: #3498db;
            color: white;
            padding: 15px;
            font-weight: bold;
            font-size: 1.1em;
        }}
        .document-content {{
            padding: 20px;
        }}
        .metadata-table {{
            width: 100%;
            border-collapse: collapse;
            margin-bottom: 15px;
        }}
        .metadata-table td {{
            padding: 8px 12px;
            border: 1px solid #ddd;
        }}
        .metadata-table td:first-child {{
            background-color: #e3f2fd;
            font-weight: bold;
            width: 150px;
        }}
        .content-section {{
            background-color: #fafafa;
            padding: 15px;
            border-radius: 5px;
            margin-top: 15px;
            border-left: 3px solid #95a5a6;
        }}
        .content-text {{
            white-space: pre-wrap;
            line-height: 1.8;
        }}
        .timestamp {{
            text-align: center;
            color: #7f8c8d;
            font-size: 0.9em;
            margin-top: 30px;
            padding-top: 20px;
            border-top: 1px solid #ecf0f1;
        }}
        .relevance-score {{
            background-color: #f39c12;
            color: white;
            padding: 2px 8px;
            border-radius: 12px;
            font-size: 0.8em;
            font-weight: bold;
        }}
    </style>
</head>
<body>
    <div class="container">
        <h1>Research Query Results</h1>
        
        <div class="query-info">
            <table>
                <tr>
                    <td>Query:</td>
                    <td>{self._escape_html(self.data.query)}</td>
                </tr>
                <tr>
                    <td>Collection:</td>
                    <td>{self._escape_html(self.data.collection_name)}</td>
                </tr>
                <tr>
                    <td>Generated:</td>
                    <td>{self.data.timestamp.strftime('%Y-%m-%d %H:%M:%S')}</td>
                </tr>
                <tr>
                    <td>Results Found:</td>
                    <td>{len(self.data.search_results)}</td>
                </tr>
            </table>
        </div>"""

        # LLM Response Section
        if self.data.llm_response:
            html_content += f"""
        <h2>AI Generated Summary</h2>
        <div class="llm-response">
            {self._escape_html(self.data.llm_response)}
        </div>"""

        # Configuration Section
        if self.data.config_info:
            html_content += """
        <h2>Configuration</h2>
        <div class="config-section">"""
            
            for key, value in self.data.config_info.items():
                if isinstance(value, dict):
                    html_content += f"<h3>{key.title()}</h3><ul>"
                    for sub_key, sub_value in value.items():
                        html_content += f"<li><strong>{sub_key}:</strong> {self._escape_html(str(sub_value))}</li>"
                    html_content += "</ul>"
                else:
                    html_content += f"<p><strong>{key.title()}:</strong> {self._escape_html(str(value))}</p>"
            
            html_content += "</div>"

        # Search Results Section
        if self.data.search_results:
            html_content += """
        <h2>Retrieved Documents</h2>"""
            
            for i, result in enumerate(self.data.search_results, 1):
                html_content += f"""
        <div class="document">
            <div class="document-header">
                Document {i}"""
                
                # Add relevance score if available
                if result.get('relevance_score') and result['relevance_score'] != 'N/A':
                    html_content += f""" <span class="relevance-score">Score: {result['relevance_score']}</span>"""
                
                html_content += """
            </div>
            <div class="document-content">"""
                
                # Metadata
                if result['metadata']:
                    filtered_metadata = filter_metadata(result['metadata'])
                    if filtered_metadata:
                        html_content += """
                <table class="metadata-table">"""
                        for key, value in filtered_metadata.items():
                            html_content += f"""
                    <tr>
                        <td>{key.title()}</td>
                        <td>{self._escape_html(str(value))}</td>
                    </tr>"""
                        html_content += """
                </table>"""
                
                # Content
                html_content += f"""
                <h3>Content</h3>
                <div class="content-section">
                    <div class="content-text">{self._escape_html(result['content'])}</div>
                </div>
            </div>
        </div>"""

        html_content += f"""
        <div class="timestamp">
            Generated on {self.data.timestamp.strftime('%B %d, %Y at %I:%M %p')}
        </div>
    </div>
</body>
</html>"""

        # Write to file
        with open(output_path, 'w', encoding='utf-8') as f:
            f.write(html_content)
    
    def _escape_html(self, text: str) -> str:
        """Escape HTML special characters"""
        if not isinstance(text, str):
            text = str(text)
        
        replacements = {
            '&': '&amp;',
            '<': '&lt;',
            '>': '&gt;',
            '"': '&quot;',
            "'": '&#x27;',
        }
        
        for char, replacement in replacements.items():
            text = text.replace(char, replacement)
        
        return text